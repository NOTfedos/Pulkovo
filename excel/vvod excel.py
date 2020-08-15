import openpyxl, os
import win32com.client as win32

#----------------------------------------------------------
class Aud(object):   
    def __init__(self, num, space, spes, priority, whitelist):
        self.num = num
        self.space = space 
        self.spes = spes #особенности
        self.priority = priority
        self.whitelist = whitelist
        
 
        

class Program(object):   
    def __init__(self, num, disp, name, spes, time, days, group, people, people_vn, teach, aud, skelet):
        self.num = num
        self.disp = disp
        self.name = name
        self.spes = spes #особенности
        self.time = time
        self.days = days
        self.group = group
        self.people = people
        self.people_vn = people_vn
        self.teach = teach
        self.aud = aud
        self.skelet = skelet
        
def whitelist_check(s):
    for c,i in enumerate(s):
        if "кроме" in i:
            s.pop(c)
    return s
def pars_lessons(ind,file_1):

    return list(map(lambda s: s.value, file_1['I'+str(ind):'DK'+str(ind)][0][::2]))


#----------------------------------------------------------
file_1 = openpyxl.load_workbook('Приложение №1.xlsx')['ДПО']

disp = {}

offset_1 = 6

index = 6

while (file_1['G'+str(index)].value != None):
    name = file_1['G'+str(index)].value
    disp[name] = pars_lessons(index, file_1)
    index += 2


del file_1

#----------------------------------------------------------
file_2 = openpyxl.load_workbook('Приложение №2.xlsx')
#---------------
file_auds = file_2['параметры аудиторий']

Audit = []
index = 2

while (file_auds['A'+str(index)].value != None):
    if 'теоретические' in file_auds['C'+str(index)].value:
        Audit.append(Aud(file_auds['A'+str(index)].value, file_auds['B'+str(index)].value, file_auds['D'+str(index)].value.split(', '), file_auds['E'+str(index)].value, whitelist_check(file_auds['F'+str(index)].value.split('; ').copy())))
    index += 1
del file_auds
#---------------

file_programs = file_2['параметры программ']

index = 2
Programs = []
last = ''

while (file_programs['A'+str(index)].value != None):
    if (last != file_programs['B'+str(index)].value) amd (file_programs['B'+str(index)].value != None):
        last = file_programs['B'+str(index)].value
    if str(file_programs['C'+str(index)].value) in os.listdir('/Приложение №3 (2)'):        

        wordDoc = docx.Document("/Приложение №3 (2)/"+str(file_programs['C'+str(index)].value))


    else:
        skelet = 0

    Programs.append(Program(file_programs['A'+str(index)].value, last, file_programs['C'+str(index)].value, file_programs['D'+str(index)].value.split(', '), file_programs['F'+str(index)].value, file_programs['I'+str(index)].value, file_programs['J'+str(index)].value, file_programs['K'+str(index)].value, file_programs['L'+str(index)].value, file_programs['M'+str(index)].value, file_programs['N'+str(index)].value, skelet))
    





table_size = len(wordDoc.tables[1].rows)
n_report = wordDoc.tables[0].rows[-1].cells[0].text
date_report = wordDoc.tables[0].rows[-1].cells[1].text

dataframe = []
for i in range(1,table_size):
    dataframe.append(wordDoc.tables[1].rows[i].cells[1].text)
dataframe = pd.DataFrame(dataframe)
print(dataframe)